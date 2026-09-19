# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1,<4"]
# ///
"""Markdown -> xlsx (one sheet per pipe table found in the document; the
heading immediately above a table, if any, names its sheet). Standalone — no
daemon/Jones import.

Usage:
    uv run md_to_xlsx.py <input.md> <output.xlsx>

A document with no pipe table still produces an openable workbook: one sheet
holding the document's paragraphs/headings as plain rows, so this script
never refuses input just because it isn't tabular (see `_shared_md.py`'s
module docstring for the same "produce an openable file" principle).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from _shared_md import parse_blocks, plain_text


def _append_text_row(ws, values: list[str]) -> None:
    """Append a row where every cell is forced to Excel's "string" type,
    no matter what the text looks like.

    Plain ``ws.append(values)`` alone hands each value to openpyxl's
    ``Cell.value`` setter, which infers the cell's type from its content —
    a value starting with "=" is bound as a live formula (``data_type =
    "f"``). Markdown table/paragraph text is untrusted (it can come from a
    scraped page or an agent-authored report), so a literal cell like
    "=1+1" or "=cmd|'/C calc'!A1" would silently become an executable
    formula in the generated workbook instead of staying literal text
    (CSV/formula injection, CWE-1236; review round-1, #3). Re-stamping
    ``data_type`` *after* the value assignment is required — the setter
    would otherwise overwrite it right back to "f" on the next value set.
    """
    ws.append(values)
    row = ws.max_row
    for col in range(1, len(values) + 1):
        ws.cell(row=row, column=col).data_type = "s"


def _sheet_title(name: str, used: set[str]) -> str:
    # Excel sheet name limits: <=31 chars, no `: \ / ? * [ ]`.
    cleaned = re.sub(r"[:\\/?*\[\]]", " ", name).strip() or "Sheet"
    cleaned = cleaned[:31]
    base, n = cleaned, 2
    while cleaned in used:
        suffix = f" ({n})"
        cleaned = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(cleaned)
    return cleaned


def build_workbook(markdown_text: str):
    import openpyxl

    blocks = parse_blocks(markdown_text)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()

    tables = [(i, payload) for i, (kind, payload) in enumerate(blocks) if kind == "table"]
    if not tables:
        ws = wb.create_sheet(_sheet_title("Document", used_titles))
        for kind, payload in blocks:
            if kind == "heading":
                _append_text_row(ws, [plain_text(payload[1])])
            elif kind == "paragraph":
                _append_text_row(ws, [plain_text(payload)])
            elif kind in ("bullet_list", "ordered_list"):
                for item in payload:
                    _append_text_row(ws, ["", plain_text(item)])
        return wb

    for idx, (block_i, rows) in enumerate(tables, start=1):
        title = f"Table {idx}"
        for j in range(block_i - 1, -1, -1):
            if blocks[j][0] == "heading":
                title = blocks[j][1][1]
                break
        ws = wb.create_sheet(_sheet_title(title, used_titles))
        for row in rows:
            _append_text_row(ws, [plain_text(cell) for cell in row])
    return wb


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(f"usage: {argv[0]} <input.md> <output.xlsx>", file=sys.stderr)
        return 2
    src, dst = Path(argv[1]), Path(argv[2])
    markdown_text = src.read_text(encoding="utf-8")
    wb = build_workbook(markdown_text)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dst))
    print(f"wrote {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
