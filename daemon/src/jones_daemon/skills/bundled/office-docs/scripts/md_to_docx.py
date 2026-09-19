# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx>=1.1,<2"]
# ///
"""Markdown -> Docx, via python-docx. Standalone — no daemon/Jones import.

Usage:
    uv run md_to_docx.py <input.md> <output.docx>

Deliberately not `uv run --with python-docx md_to_docx.py ...`: PEP 723's
inline `# /// script` block above already declares the dependency, so a bare
`uv run md_to_docx.py ...` resolves and caches it on first use — the agent
never needs to know (or pin) this skill's Python dependencies itself.
"""

from __future__ import annotations

import sys
from pathlib import Path

from _shared_md import parse_blocks, parse_inline


def _add_inline_runs(paragraph, text: str) -> None:
    for run in parse_inline(text):
        docx_run = paragraph.add_run(run.text)
        docx_run.bold = run.bold
        docx_run.italic = run.italic
        if run.code:
            docx_run.font.name = "Courier New"


def build_docx(markdown_text: str):
    import docx

    document = docx.Document()
    for kind, payload in parse_blocks(markdown_text):
        if kind == "heading":
            level, text = payload
            heading = document.add_heading(level=min(level, 9))
            _add_inline_runs(heading, text)
        elif kind == "paragraph":
            _add_inline_runs(document.add_paragraph(), payload)
        elif kind == "bullet_list":
            for item in payload:
                _add_inline_runs(document.add_paragraph(style="List Bullet"), item)
        elif kind == "ordered_list":
            for item in payload:
                _add_inline_runs(document.add_paragraph(style="List Number"), item)
        elif kind == "code":
            p = document.add_paragraph()
            run = p.add_run(payload)
            run.font.name = "Courier New"
        elif kind == "blockquote":
            p = document.add_paragraph(style="Intense Quote")
            _add_inline_runs(p, payload)
        elif kind == "table":
            rows = payload
            n_cols = max(len(r) for r in rows)
            table = document.add_table(rows=len(rows), cols=n_cols)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c in range(n_cols):
                    cell_text = row[c] if c < len(row) else ""
                    table.cell(r, c).text = "".join(x.text for x in parse_inline(cell_text))
        elif kind == "hr":
            document.add_paragraph("―" * 20)
        else:  # pragma: no cover - parse_blocks() never emits an unknown kind
            raise ValueError(f"unknown block kind: {kind}")
    return document


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(f"usage: {argv[0]} <input.md> <output.docx>", file=sys.stderr)
        return 2
    src, dst = Path(argv[1]), Path(argv[2])
    markdown_text = src.read_text(encoding="utf-8")
    document = build_docx(markdown_text)
    dst.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dst))
    print(f"wrote {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
